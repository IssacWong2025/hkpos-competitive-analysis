
#!/usr/bin/env python3
from __future__ import annotations

import argparse, csv, datetime as dt, html, io, json, os, re, sys, zipfile
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Sequence, Tuple
from urllib.parse import quote_plus, urlparse

try:
    import requests
except Exception:
    requests = None
try:
    import openpyxl
except Exception:
    openpyxl = None

KEYWORD_HINTS={"competitor","competitors","pos","hong kong","hk","domain","website","facebook","page","instagram","ig","ads","广告","投放"}
SCAN_EXTENSIONS={".csv",".md",".json",".html",".txt",".xlsx"}
COMPETITOR_FIELDS=["competitor_name","website_domain","website_url","facebook_page_url","instagram_handle","notes_source_file","confidence","missing_page_url"]
META_FIELDS=["competitor_name","advertiser_name","facebook_page_url","ad_library_url","ad_count_active","ad_id_or_archive_id","status","platforms_hint","ad_format_hint","objective_path_hint","objective_reason","message_destination_hint","landing_page_url","primary_text","headline","call_to_action","captured_at","collection_method","error_reason","manual_required_fields","manual_instructions"]
SEMRUSH_FIELDS=["competitor_name","website_domain","paid_keywords_top","paid_keywords_count","sample_ad_copies","database","units_before","units_after","captured_at"]
EN_STOPWORDS={"the","and","for","with","your","you","from","that","this","are","our","can","now","get","pos","hong","kong"}
ZH_STOPWORDS={"的","了","和","及","與","为","是","在","可","更","你","您"}

CORE_COMPETITORS = [
    {"competitor_name":"Eats365","website_domain":"eats365pos.com","website_url":"https://www.eats365pos.com/hk/"},
    {"competitor_name":"ezPOS","website_domain":"ezpos.hk","website_url":"https://www.ezpos.hk/"},
    {"competitor_name":"ROKA","website_domain":"roka.com.hk","website_url":"https://www.roka.com.hk/"},
    {"competitor_name":"OmniWe","website_domain":"omniwe.com","website_url":"https://omniwe.com/en-US/oursolutions/dining_pos"},
    {"competitor_name":"iCHEF","website_domain":"ichefpos.com","website_url":"https://www.ichefpos.com/zh-hk"},
    {"competitor_name":"DimPOS","website_domain":"dimorder.com","website_url":"https://www.dimorder.com/dimpos/"},
    {"competitor_name":"HCTC","website_domain":"hctc.com.hk","website_url":"https://posapp.hctc.com.hk/"},
    {"competitor_name":"Loyverse","website_domain":"loyverse.com","website_url":"https://loyverse.com/features"},
    {"competitor_name":"Gingersoft","website_domain":"clggroup.com.hk","website_url":"https://clggroup.com.hk/hk/fbPos"},
    {"competitor_name":"Caterlord","website_domain":"caterlord.com","website_url":"https://caterlord.com/zh-hant/"},
    {"competitor_name":"DoLA","website_domain":"dolatechnology.com","website_url":"https://www.dolatechnology.com/zh"},
]

@dataclass
class StepStat:
    step:str
    success:int=0
    failed:int=0

class StepLogger:
    def __init__(self,out_dir:Path)->None:
        out_dir.mkdir(parents=True,exist_ok=True)
        ts=dt.datetime.now().strftime("%Y%m%d_%H%M%S")
        self.path=out_dir/f"run_all_{ts}.log"
        self.latest=out_dir/"run_all.log"
        self._buf=[]
    def log(self,msg:str)->None:
        line=f"[{dt.datetime.now().isoformat(timespec='seconds')}] {msg}"
        print(line);self._buf.append(line)
    def flush(self)->None:
        body="\n".join(self._buf)+"\n"
        self.path.write_text(body,encoding="utf-8")
        self.latest.write_text(body,encoding="utf-8")

def sanitize_text(s:str)->str:
    if not s:return s
    s=re.sub(r"([?&]key=)[^&\s]+",r"\1***",s,flags=re.I)
    s=re.sub(r"(SEMRUSH_API_KEY=)[^\s]+",r"\1***",s,flags=re.I)
    return s

def now_iso()->str:return dt.datetime.now().isoformat(timespec="seconds")

def normalize_domain(v:str)->str:
    v=(v or "").strip().lower()
    if not v:return ""
    if not v.startswith("http"):v=f"https://{v}"
    try:netloc=urlparse(v).netloc.lower().split(":")[0]
    except Exception:return ""
    return netloc[4:] if netloc.startswith("www.") else netloc

def normalize_website_url(v:str)->str:
    v=(v or "").strip()
    if not v:return ""
    return v if re.match(r"^https?://",v,re.I) else f"https://{v}"

def normalize_ig_handle(v:str)->str:
    v=(v or "").strip()
    if not v:return ""
    v=re.sub(r"^https?://(www\.)?instagram\.com/","",v,flags=re.I)
    h=v.strip("/@")
    if h.lower() in {"media","context","keyframes","reel","p"}:
        return ""
    return h

def infer_name(domain:str,website_url:str,fallback:str="")->str:
    if (fallback or "").strip():return fallback.strip()
    d=domain or normalize_domain(website_url)
    return "" if not d else d.split(".")[0].replace("-"," ").title()

def confidence_level(r:Dict[str,str])->str:
    has_name,has_domain,has_fb,has_ig=bool(r.get("competitor_name")),bool(r.get("website_domain")),bool(r.get("facebook_page_url")),bool(r.get("instagram_handle"))
    if has_name and has_domain and (has_fb or has_ig):return "high"
    if (has_name and has_domain) or (has_name and has_fb):return "medium"
    return "low"

def dedupe_competitors(rows:Sequence[Dict[str,str]])->List[Dict[str,str]]:
    score={"low":1,"medium":2,"high":3};merged={};sources=defaultdict(set)
    for row in rows:
        r={k:(row.get(k,"") or "").strip() for k in COMPETITOR_FIELDS}
        r["website_domain"]=normalize_domain(r["website_domain"] or r["website_url"])
        r["website_url"]=normalize_website_url(r["website_url"] or r["website_domain"])
        r["instagram_handle"]=normalize_ig_handle(r["instagram_handle"])
        r["competitor_name"]=infer_name(r["website_domain"],r["website_url"],r["competitor_name"])
        r["confidence"]=r["confidence"] or confidence_level(r)
        r["missing_page_url"]="true" if not r["facebook_page_url"] else "false"
        key=(r["website_domain"] or r["competitor_name"]).lower().strip()
        if not key:continue
        sources[key].add(r.get("notes_source_file",""))
        if key not in merged:merged[key]=r;continue
        cur=merged[key];best=r if score.get(r["confidence"],1)>score.get(cur["confidence"],1) else cur;other=cur if best is r else r
        for f in COMPETITOR_FIELDS:
            if f!="notes_source_file" and not best.get(f) and other.get(f):best[f]=other[f]
        best["confidence"]=confidence_level(best);best["missing_page_url"]="true" if not best.get("facebook_page_url") else "false";merged[key]=best
    out=[]
    for k,r in merged.items():r["notes_source_file"]=";".join(sorted(x for x in sources[k] if x));out.append(r)
    out.sort(key=lambda x:(x.get("competitor_name","").lower(),x.get("website_domain","")))
    return out

def read_text_file(path:Path)->str:
    for enc in ("utf-8","utf-8-sig","cp936","gb18030","latin-1"):
        try:return path.read_text(encoding=enc)
        except Exception:pass
    return ""

def row_text_has_hints(t:str)->bool:
    t=(t or "").lower();return any(k in t for k in KEYWORD_HINTS)

def extract_domains(text:str)->List[str]:
    found=re.findall(r"\b(?:[a-zA-Z0-9-]+\.)+[a-zA-Z]{2,}\b",text or "")
    out=[]
    for d in found:
        d=d.lower().strip(".")
        if d.endswith((".png",".jpg",".jpeg",".js",".css",".json",".pdf")):continue
        out.append(d)
    return out

def is_valid_competitor_row(r:Dict[str,str])->bool:
    d=(r.get("website_domain","") or "").lower()
    n=(r.get("competitor_name","") or "").lower()
    if not d and not n:return False
    if n in {"www","com","support","common","i","li"}:return False
    if re.match(r"^[0-9]",n):return False
    bad_hosts=("facebook.com","instagram.com","similarweb.com","clarity.ms","gstatic.com","google.com","doubleclick.net","cloudfront.net","awswaf.com","facebook.net")
    if any(x in d for x in bad_hosts):return False
    parts=d.split(".")
    if len(parts)<2 or len(parts)>3:return False
    if not re.search(r"[a-z]",parts[-2] if len(parts)>=2 else ""):return False
    if parts[-1] not in {"com","hk","net","io","app","org","co","biz","us"}:return False
    if any(x in d for x in ("eael","elementor","addeventlistener","position","value","default","preset")):return False
    if len(n)>60:return False
    allow=("eats365","ezpos","roka","omniwe","ichef","dimpos","hctc","caterlord","dola","loyverse","gingersoft","clggroup","eposhk","favourpos","food")
    return any(k in d for k in allow) or any(k in n for k in allow)

def extract_facebook_urls(text:str)->List[str]:return re.findall(r"https?://(?:www\.)?facebook\.com/[A-Za-z0-9_.\-/]+",text or "",flags=re.I)
def extract_instagram_handles(text:str)->List[str]:
    urls=re.findall(r"https?://(?:www\.)?instagram\.com/[A-Za-z0-9_.\-/]+",text or "",flags=re.I)
    hs=[normalize_ig_handle(u) for u in urls]
    return [h for h in hs if h and h.lower() not in {"facebook","instagram"}]

def pick_best_facebook_url(urls:List[str])->str:
    cleaned=[]
    for u in urls:
        lu=u.lower()
        if any(x in lu for x in ["/ads/library","/sharer","/share.php","/plugins/","/dialog/","/2008/fbml"]):
            continue
        cleaned.append(u.rstrip("/"))
    if not cleaned:
        return ""
    # Prefer explicit page slugs over generic profile.php links.
    cleaned.sort(key=lambda x: (("profile.php" in x.lower()), len(x)))
    return cleaned[0]

def is_invalid_facebook_url(url:str)->bool:
    u=(url or "").lower()
    if not u:
        return True
    if any(x in u for x in ["/ads/library","/sharer","/share.php","/plugins/","/dialog/","/2008/fbml"]):
        return True
    if "profile.php" in u and "id=" not in u:
        return True
    return False

def extract_from_csv(path:Path)->List[Dict[str,str]]:
    data=read_text_file(path)
    if not data.strip():return []
    try:reader=csv.DictReader(io.StringIO(data))
    except Exception:return []
    rows=[]
    for item in reader:
        n={str(k).strip().lower():str(v or "").strip() for k,v in item.items()}
        if not row_text_has_hints(" ".join(n.values())):continue
        name=n.get("competitor") or n.get("competitor_name") or n.get("name") or n.get("company")
        website=n.get("website") or n.get("website_url") or n.get("url") or n.get("domain")
        domain=n.get("website_domain") or n.get("domain") or normalize_domain(website)
        fb=n.get("facebook_page_url") or n.get("facebook") or n.get("facebook_page")
        ig=n.get("instagram_handle") or n.get("instagram") or n.get("ig")
        rows.append({"competitor_name":infer_name(domain,website,name or ""),"website_domain":normalize_domain(domain),"website_url":normalize_website_url(website),"facebook_page_url":fb,"instagram_handle":normalize_ig_handle(ig),"notes_source_file":str(path),"confidence":"","missing_page_url":""})
    return rows
def extract_from_xlsx(path:Path)->List[Dict[str,str]]:
    if openpyxl is None:return []
    try:wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
    except Exception:return []
    rows=[]
    for ws in wb.worksheets:
        it=ws.iter_rows(values_only=True)
        try:headers=next(it)
        except StopIteration:continue
        if not headers:continue
        h=[str(x or "").strip().lower() for x in headers]
        for vals in it:
            cells=[str(v or "").strip() for v in vals]
            if not any(cells):continue
            item={h[i]:cells[i] for i in range(min(len(h),len(cells)))}
            if not row_text_has_hints(" ".join(item.values())):continue
            name=item.get("competitor") or item.get("competitor_name") or item.get("name") or item.get("company")
            website=item.get("website") or item.get("website_url") or item.get("url") or item.get("domain")
            domain=item.get("website_domain") or item.get("domain") or normalize_domain(website)
            fb=item.get("facebook_page_url") or item.get("facebook") or item.get("facebook_page")
            ig=item.get("instagram_handle") or item.get("instagram") or item.get("ig")
            rows.append({"competitor_name":infer_name(domain,website,name or ""),"website_domain":normalize_domain(domain),"website_url":normalize_website_url(website),"facebook_page_url":fb,"instagram_handle":normalize_ig_handle(ig),"notes_source_file":str(path),"confidence":"","missing_page_url":""})
    return rows

def extract_from_text_like(path:Path)->List[Dict[str,str]]:
    t=read_text_file(path)
    if not t or not row_text_has_hints(t):return []
    domains,fb_urls,ig_handles=extract_domains(t),extract_facebook_urls(t),extract_instagram_handles(t)
    default_fb=fb_urls[0] if fb_urls else "";default_ig=ig_handles[0] if ig_handles else ""
    return [{"competitor_name":infer_name(d,"",""),"website_domain":normalize_domain(d),"website_url":normalize_website_url(d),"facebook_page_url":default_fb,"instagram_handle":default_ig,"notes_source_file":str(path),"confidence":"","missing_page_url":""} for d in domains]

def extract_competitors_from_sources(scan_root:Path,logger:StepLogger)->List[Dict[str,str]]:
    raw=[];scanned=0
    skip_dir_tokens={"node_modules",".git","__pycache__","component_crx_cache","input"}
    for root,dirs,files in os.walk(scan_root, topdown=True):
        root_str=str(root).lower()
        dirs[:] = [d for d in dirs if d.lower() not in skip_dir_tokens and "extracted_hk_pos_competitive_analysis" not in (root_str + "\\" + d.lower())]
        if "input\\extracted_hk_pos_competitive_analysis\\hk-pos-competitive-analysis\\input\\extracted_hk_pos_competitive_analysis" in root_str:
            continue
        for fn in files:
            try:
                p=Path(root) / fn
                if p.suffix.lower() not in SCAN_EXTENSIONS:
                    continue
                scanned+=1
                if p.suffix.lower()==".csv":raw.extend(extract_from_csv(p))
                elif p.suffix.lower()==".xlsx":raw.extend(extract_from_xlsx(p))
                else:raw.extend(extract_from_text_like(p))
            except Exception as exc:
                logger.log(f"WARN extract failed: {Path(root) / fn} -> {exc}")
    for c in CORE_COMPETITORS:
        raw.append({"competitor_name":c["competitor_name"],"website_domain":c["website_domain"],"website_url":c["website_url"],"facebook_page_url":"","instagram_handle":"","notes_source_file":"seed_core","confidence":"low","missing_page_url":"true"})
    out=[x for x in dedupe_competitors(raw) if is_valid_competitor_row(x)]
    logger.log(f"Scanned {scanned} files, extracted {len(raw)} rows, filtered/deduped to {len(out)} competitors")
    return out

def enrich_social_links(competitors:List[Dict[str,str]],timeout_sec:int,logger:StepLogger)->List[Dict[str,str]]:
    if requests is None:return competitors
    for r in competitors:
        if is_invalid_facebook_url(r.get("facebook_page_url","")):
            r["facebook_page_url"]=""
        r["instagram_handle"]=normalize_ig_handle(r.get("instagram_handle",""))
        if r.get("facebook_page_url") and r.get("instagram_handle"):continue
        url=r.get("website_url","")
        if not url:continue
        try:
            resp=requests.get(url,headers={"User-Agent":"Mozilla/5.0","Accept-Language":"en-US,en;q=0.9"},timeout=timeout_sec)
            if resp.status_code!=200:continue
            body=resp.text
            if not r.get("facebook_page_url"):
                fb=extract_facebook_urls(body)
                best_fb=pick_best_facebook_url(fb)
                if best_fb:r["facebook_page_url"]=best_fb
            if not r.get("instagram_handle"):
                ig=extract_instagram_handles(body)
                if ig:r["instagram_handle"]=ig[0]
            # Filter noisy FB placeholders commonly found in widgets.
            fbv=(r.get("facebook_page_url","") or "").lower()
            if is_invalid_facebook_url(fbv):
                r["facebook_page_url"]=""
            r["missing_page_url"]="true" if not (r.get("facebook_page_url") or r.get("instagram_handle")) else "false"
            r["confidence"]=confidence_level(r)
        except Exception:
            continue
    logger.log(f"Social enrichment done, with_social_url={sum(1 for c in competitors if c.get('facebook_page_url') or c.get('instagram_handle'))}")
    # Curated fallback for core competitors when official site link is not directly exposed.
    curated_fb={
        "ichef":"https://www.facebook.com/iCHEF/",
        "loyverse":"https://www.facebook.com/loyversepos/",
        "omniwe":"https://www.facebook.com/omniwe",
    }
    for r in competitors:
        if r.get("facebook_page_url"):
            continue
        key=(r.get("competitor_name","")+r.get("website_domain","")).lower()
        for token,url in curated_fb.items():
            if token in key:
                r["facebook_page_url"]=url
                r["missing_page_url"]="false"
                break
    return competitors

def normalize_baseline_competitors(rows:List[Dict[str,str]])->List[Dict[str,str]]:
    canon={}
    for c in CORE_COMPETITORS:
        canon[c["competitor_name"].lower()]=c
        canon[normalize_domain(c["website_domain"]).split(".")[0]]=c
    out=[]
    for r in rows:
        k=(r.get("competitor_name","")+r.get("website_domain","")).lower()
        for token,patch in canon.items():
            if token in k:
                r.update({**patch,**r})
                for pk,pv in patch.items():
                    if pk in ("competitor_name","website_domain","website_url"):r[pk]=pv
                break
        r["missing_page_url"]="true" if not (r.get("facebook_page_url") or r.get("instagram_handle")) else "false"
        r["confidence"]=confidence_level(r)
        out.append(r)
    merged=dedupe_competitors(out)
    # Guarantee core competitors exist in master list.
    existing={normalize_domain(x.get("website_domain","") or x.get("website_url","")) for x in merged}
    for c in CORE_COMPETITORS:
        dom=normalize_domain(c["website_domain"])
        if dom in existing:
            continue
        merged.append({
            "competitor_name":c["competitor_name"],
            "website_domain":dom,
            "website_url":c["website_url"],
            "facebook_page_url":"",
            "instagram_handle":"",
            "notes_source_file":"seed_core",
            "confidence":"medium",
            "missing_page_url":"true",
        })
    return dedupe_competitors(merged)

def write_csv(path:Path,fields:Sequence[str],rows:Sequence[Dict[str,str]])->None:
    path.parent.mkdir(parents=True,exist_ok=True)
    with path.open("w",newline="",encoding="utf-8") as f:
        w=csv.DictWriter(f,fieldnames=fields);w.writeheader()
        for r in rows:w.writerow({k:r.get(k,"") for k in fields})

def read_csv_rows(path:Path)->List[Dict[str,str]]:
    if not path.exists():
        return []
    with path.open("r",encoding="utf-8",newline="") as f:
        return list(csv.DictReader(f))

def build_ad_library_url(facebook_page_url:str,competitor_name:str)->str:
    slug=""
    if facebook_page_url:
        try:
            parts=[x for x in urlparse(facebook_page_url).path.split("/") if x];slug=parts[0] if parts else ""
        except Exception:slug=""
    q=(slug or competitor_name or "").strip()
    return "https://www.facebook.com/ads/library/?active_status=active&ad_type=all&country=HK&is_targeted_country=false&media_type=all&search_type=keyword&q="+quote_plus(q)

def parse_meta_ad_library_html(body:str)->Dict[str,str]:
    p={}
    if not body:return p
    m=re.search(r'"result_count"\s*:\s*(\d+)',body)
    if m:p["ad_count_active"]=m.group(1)
    a=re.search(r'"ad_archive_id"\s*:\s*"?(\d+)"?',body)
    if a:p["ad_id_or_archive_id"]=a.group(1)
    cta=re.search(r'"call_to_action_type"\s*:\s*"([A-Z_]+)"',body)
    if cta:p["call_to_action"]=cta.group(1).lower()
    link=re.search(r'https?://[^"\s]+',body)
    if link:p["landing_page_url"]=html.unescape(link.group(0))
    b=re.search(r'"body"\s*:\s*\{"text"\s*:\s*"([^"]+)"',body)
    if b:p["primary_text"]=html.unescape(b.group(1))
    h=re.search(r'"title"\s*:\s*"([^"]+)"',body)
    if h:p["headline"]=html.unescape(h.group(1))
    return p

def infer_objective_and_destination(r:Dict[str,str])->Tuple[str,str,str]:
    t=" ".join([r.get("primary_text",""),r.get("headline",""),r.get("call_to_action",""),r.get("landing_page_url","")]).lower()
    if any(k in t for k in ["whatsapp","wa.me","send_message","message","messenger","dm"]):
        dest="whatsapp" if ("whatsapp" in t or "wa.me" in t) else "messenger"
        return "click_to_message",dest,"Matched messaging pattern (wa.me/whatsapp/messenger/send_message)."
    if any(k in t for k in ["instant form","lead form","lead","contact form","sign up","register","demo","book"]):
        return "lead_form","website","Matched lead-generation pattern (instant form/lead form/register/demo/book)."
    if any(k in t for k in ["http://","https://","www."]):
        return "website","website","Matched external-link pattern (website URL in creative fields)."
    return "unknown","unknown","Insufficient signals in parsed fields; needs manual verification from Ad Library card."

def infer_format_hint(r:Dict[str,str])->str:
    t=" ".join([r.get("primary_text",""),r.get("headline","")]).lower()
    if "carousel" in t:return "carousel"
    if "video" in t:return "video"
    return "image" if t else "unknown"

def collect_meta_ads(competitors:Sequence[Dict[str,str]],timeout_sec:int,logger:StepLogger)->Tuple[List[Dict[str,str]],List[Dict[str,str]]]:
    rows=[];todos=[]
    if requests is None:logger.log("requests not installed; all Meta rows set to manual_needed")
    for c in competitors:
        name,fb_url=c.get("competitor_name",""),c.get("facebook_page_url","")
        base={"competitor_name":name,"advertiser_name":name,"facebook_page_url":fb_url,"ad_library_url":build_ad_library_url(fb_url,name),"ad_count_active":"","ad_id_or_archive_id":"","status":"unknown","platforms_hint":"unknown","ad_format_hint":"unknown","objective_path_hint":"unknown","objective_reason":"Insufficient signals; manual review required.","message_destination_hint":"unknown","landing_page_url":"","primary_text":"","headline":"","call_to_action":"","captured_at":now_iso(),"collection_method":"manual_needed","error_reason":"","manual_required_fields":"call_to_action,landing_page_url_or_message_destination","manual_instructions":"Open ad_library_url and fill CTA type + landing page URL or WhatsApp/Messenger destination."}
        if not fb_url and not c.get("instagram_handle",""):
            base["error_reason"]="missing_facebook_or_instagram"
            rows.append(base);todos.append(base.copy());continue
        if not fb_url:
            base["error_reason"]="missing_facebook_page_url"
            rows.append(base);todos.append(base.copy());continue
        if requests is None:base["error_reason"]="requests_not_available";rows.append(base);todos.append(base.copy());continue
        try:
            resp=requests.get(base["ad_library_url"],headers={"User-Agent":"Mozilla/5.0","Accept-Language":"en-US,en;q=0.9"},timeout=timeout_sec)
            if resp.status_code!=200:
                base["error_reason"]=f"http_{resp.status_code}"
                if resp.status_code in (401,403):
                    base["status"]="unknown_blocked"
                    base["objective_reason"]="Access blocked by Meta Ad Library (http_403/http_401), requires manual confirmation."
                rows.append(base);todos.append(base.copy());continue
            parsed=parse_meta_ad_library_html(resp.text);merged=base.copy();merged.update(parsed)
            o,d,reason=infer_objective_and_destination(merged);merged["objective_path_hint"]=o;merged["message_destination_hint"]=d;merged["objective_reason"]=reason;merged["ad_format_hint"]=infer_format_hint(merged);merged["platforms_hint"]="facebook/instagram"
            if parsed:
                merged["collection_method"]="web_auto"
                merged["manual_required_fields"]=""
                merged["manual_instructions"]=""
                if merged.get("ad_count_active"):
                    try:merged["status"]="active" if int(merged["ad_count_active"])>0 else "unknown"
                    except Exception:merged["status"]="unknown"
                rows.append(merged)
            else:
                merged["collection_method"]="manual_needed";merged["error_reason"]="dynamic_render_or_no_parseable_fields";rows.append(merged);todos.append(merged.copy())
        except Exception as exc:
            base["error_reason"]=f"request_error:{exc.__class__.__name__}"
            base["status"]="unknown_blocked"
            base["objective_reason"]="Request failed while querying Ad Library; requires manual confirmation."
            rows.append(base);todos.append(base.copy())
    return rows,todos

def merge_manual_todo_fields(new_todos:List[Dict[str,str]],existing_todos:List[Dict[str,str]])->List[Dict[str,str]]:
    if not existing_todos:
        return new_todos
    if not new_todos:
        return existing_todos
    idx={(r.get("competitor_name",""),r.get("ad_library_url","")):r for r in existing_todos}
    idx_by_comp={r.get("competitor_name",""):r for r in existing_todos}
    keep_fields=["ad_count_active","ad_id_or_archive_id","status","objective_path_hint","objective_reason","message_destination_hint","landing_page_url","primary_text","headline","call_to_action","manual_required_fields","manual_instructions"]
    def should_keep(field:str,val:str)->bool:
        v=(val or "").strip()
        if not v:
            return False
        if field=="status":
            return v.lower() not in {"unknown","unknown_blocked"}
        if field=="objective_path_hint":
            return v.lower() not in {"unknown"}
        if field=="objective_reason":
            return "insufficient signals" not in v.lower()
        return True
    out=[];seen=set()
    for row in new_todos:
        key=(row.get("competitor_name",""),row.get("ad_library_url",""))
        seen.add(key)
        old=idx.get(key,{})
        if not old:
            old=idx_by_comp.get(row.get("competitor_name",""),{})
        merged=row.copy()
        for f in keep_fields:
            if should_keep(f, old.get(f,"")):
                merged[f]=old.get(f)
        out.append(merged)
    # Preserve existing manual rows even if current auto-run returned no matching todo row.
    for old in existing_todos:
        key=(old.get("competitor_name",""),old.get("ad_library_url",""))
        if key in seen:
            continue
        if old.get("competitor_name",""):
            out.append(old.copy())
    return out

def apply_manual_overrides(meta_rows:List[Dict[str,str]],todo_rows:List[Dict[str,str]])->List[Dict[str,str]]:
    idx={(r.get("competitor_name",""),r.get("ad_library_url","")):r for r in todo_rows}
    idx_by_comp={r.get("competitor_name",""):r for r in todo_rows}
    writable=["ad_count_active","ad_id_or_archive_id","status","objective_path_hint","objective_reason","message_destination_hint","landing_page_url","primary_text","headline","call_to_action"]
    out=[]
    for row in meta_rows:
        key=(row.get("competitor_name",""),row.get("ad_library_url",""))
        todo=idx.get(key,{})
        if not todo:
            todo=idx_by_comp.get(row.get("competitor_name",""),{})
        merged=row.copy()
        for f in writable:
            if todo.get(f):
                merged[f]=todo.get(f)
        if merged.get("ad_count_active","").isdigit() and int(merged["ad_count_active"])>0 and merged.get("status","") in {"unknown","unknown_blocked",""}:
            merged["status"]="active"
        # If manual Ad Library evidence was provided, treat as active unless explicitly set otherwise.
        has_manual_evidence=any([
            bool((merged.get("ad_id_or_archive_id","") or "").strip()),
            bool((merged.get("primary_text","") or "").strip()),
            bool((merged.get("headline","") or "").strip()),
            bool((merged.get("call_to_action","") or "").strip()),
            bool((merged.get("landing_page_url","") or "").strip()),
        ])
        if has_manual_evidence and merged.get("status","").lower() in {"", "unknown", "unknown_blocked"}:
            merged["status"]="active"
            if "manual" not in (merged.get("objective_reason","").lower()):
                merged["objective_reason"]=(merged.get("objective_reason","").strip()+" Manual Ad Library fields populated (copy/CTA/landing/ad_id), marked as active evidence.").strip()
        # Re-infer objective when still unknown but manual fields now exist.
        if (merged.get("objective_path_hint","") or "").lower() in {"", "unknown"}:
            o,d,reason=infer_objective_and_destination(merged)
            merged["objective_path_hint"]=o
            merged["message_destination_hint"]=d
            if reason:
                merged["objective_reason"]=reason
        out.append(merged)
    return out

def build_manual_todo_rows(meta_rows:Sequence[Dict[str,str]])->List[Dict[str,str]]:
    out=[]
    for r in meta_rows:
        need=[]
        if not (r.get("call_to_action","") or "").strip():
            need.append("call_to_action")
        if not (r.get("landing_page_url","") or "").strip() and (r.get("message_destination_hint","unknown") or "unknown")=="unknown":
            need.append("landing_page_url_or_message_destination")
        if (r.get("objective_path_hint","unknown") or "unknown")=="unknown":
            need.append("objective_path_hint")
        if not need and (r.get("status","") or "").lower() in {"active"}:
            continue
        row=r.copy()
        row["manual_required_fields"]=",".join(need) if need else "status_confirmation"
        row["manual_instructions"]="Open ad_library_url and fill CTA + destination (landing URL or WhatsApp/Messenger) and objective hint."
        out.append(row)
    return out

def tokenize_mixed_text(t:str)->List[str]:
    en=[w.lower() for w in re.findall(r"[A-Za-z][A-Za-z0-9_']{1,}",t or "") if w.lower() not in EN_STOPWORDS]
    zh=[]
    for seq in re.findall(r"[\u4e00-\u9fff]{2,}",t or ""):
        if seq in ZH_STOPWORDS:continue
        zh.append(seq)
        if len(seq)>=4:
            for i in range(len(seq)-1):
                g=seq[i:i+2]
                if g not in ZH_STOPWORDS:zh.append(g)
    return en+zh

def extract_copy_keywords(texts:Sequence[str],top_k_words:int=10,top_k_phrases:int=5)->Tuple[List[Tuple[str,int]],List[Tuple[str,int]]]:
    wc,pc=Counter(),Counter()
    for t in texts:
        c=(t or "").strip()
        if not c:continue
        toks=tokenize_mixed_text(c);wc.update(toks)
        if len(toks)>=2:
            for n in (2,3):
                for i in range(0,len(toks)-n+1):
                    p=" ".join(toks[i:i+n]).strip()
                    if len(p)>=4:pc[p]+=1
        pc.update(re.findall(r"[\u4e00-\u9fff]{3,12}",c))
    return wc.most_common(top_k_words),pc.most_common(top_k_phrases)
TAG_RULES = {
    "无合约":["no contract","contract-free","cancel anytime","無合約","免合約","无合约"],
    "扫码点单":["qr","scan","qrcode","qr code","掃碼","扫码","點餐","点单"],
    "外卖整合":["delivery","foodpanda","ubereats","deliveroo","外賣","外卖","外送","整合"],
    "上手快":["easy","quick setup","minutes","onboard","fast","快速","上手","簡單","简单"],
    "硬件兼容":["ipad","tablet","printer","hardware","硬件","設備","设备","兼容"],
    "支付收银":["payment","checkout","cashier","payme","octopus","credit card","收銀","收银","支付"],
    "提升效率":["efficiency","streamline","faster","turnover","效率","翻台","省時","省时"],
    "促销优惠":["discount","offer","free trial","優惠","优惠","折扣","試用","试用"],
    "私信咨询":["send_message","message us","whatsapp","messenger","私訊","私信","查詢","咨询"],
}

def normalize_copy_text(text:str)->str:
    s=(text or "").strip()
    if not s:
        return ""
    # Decode escaped unicode safely when text is pasted as \uXXXX sequences.
    if "\\u" in s:
        def _decode_u(m:re.Match)->str:
            try:
                return chr(int(m.group(1),16))
            except Exception:
                return m.group(0)
        s=re.sub(r"\\u([0-9a-fA-F]{4})",_decode_u,s)
    # Remove template placeholders that do not carry proposition meaning.
    s = re.sub(r"\{\{[^}]+\}\}", " ", s)
    s = s.replace("\\n", " ").replace("\n", " ")
    return re.sub(r"\s+", " ", s).strip()

def classify_proposition_tags(text:str, objective_hint:str="", cta:str="", destination:str="")->Tuple[str,str,str]:
    norm=normalize_copy_text(text)
    t=norm.lower()
    scores=Counter()
    reasons=defaultdict(list)
    for tag,patterns in TAG_RULES.items():
        for p in patterns:
            if p.lower() in t:
                scores[tag]+=1
                reasons[tag].append(p)
    # Objective/CTA fallback signals.
    obj=(objective_hint or "").lower()
    cta_l=(cta or "").lower()
    dest=(destination or "").lower()
    if obj=="click_to_message" or any(x in cta_l for x in ["message","whatsapp","send"]):
        scores["私信咨询"] += 2
        reasons["私信咨询"].append("objective/cta message signal")
    if obj=="website" and any(x in t for x in ["book","learn more","了解","demo","官网","website"]):
        scores["上手快"] += 1
        reasons["上手快"].append("website exploration signal")
    if dest=="website":
        scores["网站引流"] += 1
        reasons["网站引流"].append("message_destination_hint=website")
    if not scores:
        return "待人工判定","","No tag keyword matched in primary_text/headline/CTA."
    ranked=[x[0] for x in scores.most_common(2)]
    r1="; ".join(reasons[ranked[0]]) if ranked else ""
    r2="; ".join(reasons[ranked[1]]) if len(ranked)>1 else ""
    reason=f"primary={ranked[0]} via [{r1}]"
    if len(ranked)>1:
        reason+=f"; secondary={ranked[1]} via [{r2}]"
    return ranked[0], (ranked[1] if len(ranked)>1 else ""), reason

def build_meta_keyword_rows(meta_rows:Sequence[Dict[str,str]])->List[Dict[str,str]]:
    out=[];summary=defaultdict(Counter)
    for r in meta_rows:
        text=" ".join([r.get("primary_text",""),r.get("headline",""),r.get("call_to_action","")]).strip()
        p,s,reason=classify_proposition_tags(
            text,
            objective_hint=r.get("objective_path_hint",""),
            cta=r.get("call_to_action",""),
            destination=r.get("message_destination_hint",""),
        )
        comp=r.get("competitor_name","")
        if p:summary[comp][p]+=1
        if s:summary[comp][s]+=1
        out.append({
            "row_type":"ad_tag",
            "competitor_name":comp,
            "ad_id_or_archive_id":r.get("ad_id_or_archive_id",""),
            "ad_library_url":r.get("ad_library_url",""),
            "label_primary":p,
            "label_secondary":s,
            "label_reason":reason,
            "count":"1",
        })
    for comp,counter in summary.items():
        for label,count in counter.most_common():
            out.append({
                "row_type":"competitor_summary",
                "competitor_name":comp,
                "ad_id_or_archive_id":"",
                "ad_library_url":"",
                "label_primary":label,
                "label_secondary":"",
                "label_reason":"aggregated tag count",
                "count":str(count),
            })
    return out

def build_meta_type_distribution(meta_rows:Sequence[Dict[str,str]])->Dict[str,int]:
    by_comp=defaultdict(list)
    for r in meta_rows:by_comp[r.get("competitor_name","")].append(r.get("objective_path_hint","unknown"))
    counts={"click_to_message":0,"lead_form":0,"website":0,"unknown":0}
    for _,objs in by_comp.items():
        f=[o for o in objs if o and o!="unknown"];chosen=Counter(f).most_common(1)[0][0] if f else "unknown";counts[chosen]=counts.get(chosen,0)+1
    return counts

def semrush_units(api_key:str)->str:
    if not requests:return ""
    try:
        r=requests.get("https://api.semrush.com/",params={"type":"api_units","key":api_key},timeout=20)
        return r.text.strip() if r.status_code==200 else ""
    except Exception:
        return ""

def semrush_units_info(api_key:str)->Tuple[str,str]:
    if not requests:
        return "","requests_not_available"
    body,err=semrush_request({"type":"api_units","key":api_key},timeout=20,retries=1)
    if err:
        return "",err
    if not body:
        return "","empty_units_response"
    return body.strip(),""

def semrush_request(params:Dict[str,str],timeout:int=30,retries:int=2)->Tuple[str,str]:
    if not requests:return "","requests_not_available"
    last=""
    for _ in range(retries+1):
        try:
            r=requests.get("https://api.semrush.com/",params=params,timeout=timeout)
            if r.status_code!=200:
                last=f"http_{r.status_code}"
                continue
            return r.text,""
        except Exception as exc:
            last=f"request_error:{exc.__class__.__name__}"
    return "",last or "request_failed"

def semrush_paid_keywords(api_key:str,domain:str,database:str,display_limit:int)->Tuple[List[Dict[str,str]],str]:
    p={"type":"domain_adwords","key":api_key,"domain":domain,"database":database,"display_limit":str(display_limit),"export_escape":"1","export_columns":"Ph,Po,Pp,Nq,Cp"}
    b,err=semrush_request(p,timeout=30,retries=2)
    if err:return [],err
    if b.startswith("ERROR"):return [],b.strip().replace("\n"," ")
    lines=[ln.strip() for ln in b.splitlines() if ln.strip()]
    if not lines:return [],"empty_response"
    first=[x.strip() for x in lines[0].split(";")]
    expected={"Ph","Po","Pp","Nq","Cp"}
    if expected.intersection(set(first)):
        h=first
        data_lines=lines[1:]
    else:
        # Some accounts return data-only rows without header; map by export_columns order.
        h=["Ph","Po","Pp","Nq","Cp"]
        data_lines=lines
    out=[]
    for ln in data_lines:
        parts=[x.strip() for x in ln.split(";")]
        row={h[i]:parts[i] if i<len(parts) else "" for i in range(len(h))}
        kw=(row.get("Ph","") or "").strip().strip('"').strip("'").strip()
        pos=(row.get("Po","") or "").strip().strip('"').strip("'").strip()
        if kw.lower()=="keyword" and pos.lower()=="position":
            continue
        row["Ph"]=kw
        row["Po"]=pos
        row["Pp"]=(row.get("Pp","") or "").strip().strip('"').strip("'").strip()
        row["Nq"]=(row.get("Nq","") or "").strip().strip('"').strip("'").strip()
        row["Cp"]=(row.get("Cp","") or "").strip().strip('"').strip("'").strip()
        out.append(row)
    return out,""

def semrush_sample_ads(api_key:str,domain:str,database:str,display_limit:int=20)->Tuple[List[str],str]:
    for t in ["domain_adwords_adwords","domain_adwords_ads","domain_adwords_unique"]:
        p={"type":t,"key":api_key,"domain":domain,"database":database,"display_limit":str(display_limit)}
        b,err=semrush_request(p,timeout=30,retries=1)
        if err or b.startswith("ERROR"):continue
        lines=[ln.strip() for ln in b.splitlines() if ln.strip()]
        if len(lines)>=2:return lines[1:min(len(lines),21)],""
    return [],"ads_copy_not_supported_or_not_available"

def collect_semrush_signals(competitors:Sequence[Dict[str,str]],api_key:str,database:str,display_limit:int,dry_run_count:int,logger:StepLogger)->List[Dict[str,str]]:
    domains=[]
    for c in competitors:
        d=normalize_domain(c.get("website_domain","") or c.get("website_url",""))
        if d:domains.append((c.get("competitor_name",""),d))
    dedup=[];seen=set()
    for c,d in domains:
        if d in seen:continue
        seen.add(d);dedup.append((c,d))
    logger.log(f"Quick reconnaissance: competitors={len(competitors)}, with_facebook_page_url={sum(1 for c in competitors if c.get('facebook_page_url'))}, semrush_domains={len(dedup)}, display_limit={display_limit}")
    rows=[]
    for phase,items in [("dry-run",dedup[:dry_run_count]),("batch",dedup[dry_run_count:])]:
        logger.log(f"Semrush phase={phase}, domains={len(items)}")
        for comp,domain in items:
            before,before_err=semrush_units_info(api_key);kws,kw_err=semrush_paid_keywords(api_key,domain,database,display_limit);copies,ads_err=semrush_sample_ads(api_key,domain,database,20);after,after_err=semrush_units_info(api_key)
            top=[{"keyword":k.get("Ph",""),"position":k.get("Po",""),"cpc":k.get("Cp",""),"traffic_percent":k.get("Pp","")} for k in kws[:20]]
            sample=copies[:20]
            if kw_err:sample.append(f"[paid_keywords_error] {kw_err}")
            if ads_err:sample.append(f"[ads_copy_note] {ads_err}")
            if not kws and not kw_err:
                sample.append("[paid_keywords_note] no_paid_keywords_detected")
            units_before_val = before if before else ("unavailable:" + (before_err or "unknown"))
            units_after_val = after if after else ("unavailable:" + (after_err or "unknown"))
            rows.append({"competitor_name":comp,"website_domain":domain,"paid_keywords_top":json.dumps(top,ensure_ascii=False),"paid_keywords_count":str(len(kws)) if kws else "","sample_ad_copies":" || ".join(sample),"database":database,"units_before":units_before_val,"units_after":units_after_val,"captured_at":now_iso()})
            logger.log(f"Semrush {phase}: {domain}, keywords={len(kws)}, units_before={(before if before else 'unavailable')}, units_after={(after if after else 'unavailable')}")
    return rows

def build_semrush_placeholder_rows(competitors:Sequence[Dict[str,str]],database:str,note:str)->List[Dict[str,str]]:
    rows=[];seen=set()
    for c in competitors:
        d=normalize_domain(c.get("website_domain","") or c.get("website_url",""))
        if not d or d in seen:continue
        seen.add(d)
        rows.append({"competitor_name":c.get("competitor_name",""),"website_domain":d,"paid_keywords_top":"[]","paid_keywords_count":"","sample_ad_copies":note,"database":database,"units_before":"unavailable:missing_api_key","units_after":"unavailable:missing_api_key","captured_at":now_iso()})
    return rows

def summarize_meta_keywords(rows:Sequence[Dict[str,str]])->Tuple[List[Tuple[str,int,str]],List[Tuple[str,int,str]]]:
    kwc,phc=Counter(),Counter();kwo,pho={},{}
    for r in rows:
        if r.get("row_type")!="competitor_summary":
            continue
        item=r.get("label_primary","");cnt=int(r.get("count","0") or 0);comp=r.get("competitor_name","")
        if item:
            kwc[item]+=cnt
            kwo.setdefault(item,comp)
        sec=r.get("label_secondary","")
        if sec:
            phc[sec]+=cnt
            pho.setdefault(sec,comp)
    return [(k,c,kwo.get(k,"")) for k,c in kwc.most_common(15)],[(k,c,pho.get(k,"")) for k,c in phc.most_common(10)]

def aggregate_google_intent(rows:Sequence[Dict[str,str]])->Tuple[List[Tuple[str,int]],List[Tuple[str,int]]]:
    kwc=Counter();ds=defaultdict(int)
    for r in rows:
        d=r.get("website_domain","");ds[d]+=int(r.get("paid_keywords_count","0") or 0)
        try:arr=json.loads(r.get("paid_keywords_top","") or "[]")
        except Exception:arr=[]
        for o in arr:
            kw=(o.get("keyword") or "").strip().strip('"').strip("'").strip().lower()
            if kw and kw not in {"keyword"}:
                kwc[kw]+=1
    return kwc.most_common(20),sorted(ds.items(),key=lambda x:x[1],reverse=True)[:5]

def build_ads_snapshot(path:Path,meta_rows:Sequence[Dict[str,str]],kw_rows:Sequence[Dict[str,str]],sem_rows:Sequence[Dict[str,str]])->None:
    active=[];unknown=[];likely=[]
    for r in meta_rows:
        item={
            "competitor_name":r.get("competitor_name",""),
            "status":r.get("status",""),
            "ad_library_url":r.get("ad_library_url",""),
            "ad_count_active":r.get("ad_count_active",""),
            "objective_path_hint":r.get("objective_path_hint","unknown"),
            "objective_reason":r.get("objective_reason",""),
            "message_destination_hint":r.get("message_destination_hint","unknown"),
        }
        st=(r.get("status","") or "").lower()
        if st=="active":
            active.append(item)
        elif st=="unknown_blocked":
            likely.append(item)
        else:
            unknown.append(item)
    type_distribution=build_meta_type_distribution(meta_rows)
    tag_rows=[r for r in kw_rows if r.get("row_type")=="competitor_summary"]
    top_tags=[]
    for r in tag_rows:
        lbl=(r.get("label_primary","") or "").strip()
        if not lbl:
            continue
        try:
            cnt=int(float(r.get("count","0") or 0))
        except Exception:
            cnt=0
        top_tags.append({"label":lbl,"count":cnt,"competitor_name":r.get("competitor_name","")})
    top_tags=sorted(top_tags,key=lambda x:x["count"],reverse=True)[:15]
    g20,top_domains=aggregate_google_intent(sem_rows)
    snapshot={
        "generated_at":now_iso(),
        "meta":{
            "active":active,
            "likely_active":likely,
            "unknown":unknown,
            "type_distribution":type_distribution,
            "top_tags":top_tags,
        },
        "google":{
            "top_intent_keywords":[{"keyword":k,"count":c} for k,c in g20],
            "top_domains":[{"website_domain":d,"coverage":c} for d,c in top_domains],
            "has_data":bool(sem_rows),
        }
    }
    path.parent.mkdir(parents=True,exist_ok=True)
    path.write_text(json.dumps(snapshot,ensure_ascii=False,indent=2),encoding="utf-8")

def generate_report(path:Path,competitors:Sequence[Dict[str,str]],meta_rows:Sequence[Dict[str,str]],kw_rows:Sequence[Dict[str,str]],sem_rows:Sequence[Dict[str,str]])->None:
    active=[r for r in meta_rows if r.get("status")=="active"]
    likely=[r for r in meta_rows if r.get("status") in {"unknown_blocked"}]
    unknown=[r for r in meta_rows if r.get("status") not in {"active","unknown_blocked"}]
    dist=build_meta_type_distribution(meta_rows)
    top_kw,top_ph=summarize_meta_keywords(kw_rows);g20,top_domains=aggregate_google_intent(sem_rows);reps=[f"- {r.get('competitor_name','')}: {r.get('ad_library_url','')}" for r in meta_rows[:3]]
    lines=["# HK Competitor Ads Summary","",f"Generated at: {now_iso()}","","## 1) 谁在投（Meta）","","### Active (confirmed)"]
    lines.extend([f"- {r.get('competitor_name','')} | {r.get('ad_library_url','')}" for r in active] or ["- None confirmed yet (auto/manual fields not sufficient)."])
    lines.extend(["","### Likely Active (blocked/needs manual confirm)"])
    lines.extend([f"- {r.get('competitor_name','')} | {r.get('ad_library_url','')} | reason: {r.get('error_reason','') or r.get('objective_reason','')}" for r in likely] or ["- None"])
    lines.extend(["","### Unknown / Manual check needed"])
    lines.extend([f"- {r.get('competitor_name','')} | {r.get('ad_library_url','')}" for r in unknown] or ["- None"])
    missing=[c for c in competitors if c.get("missing_page_url")=="true"]
    lines.extend(["","### Missing Facebook Page URL (manual completion queue)"])
    lines.extend([f"- {c.get('competitor_name','')} | {c.get('website_domain','')}" for c in missing] or ["- None"])
    lines.extend(["","## 2) 投什么广告类型（Meta）","",f"- Click-to-Message: {dist.get('click_to_message',0)} competitors",f"- Lead Form: {dist.get('lead_form',0)} competitors",f"- Website: {dist.get('website',0)} competitors","","Representative cases:"])
    lines.extend(reps or ["- No cases extracted"])
    lines.extend(["","## 3) Meta 文案关键词与主张","","### Top 15 主张标签（Primary）"])
    lines.extend([f"- {kw} ({cnt}) | competitor: {comp}" for kw,cnt,comp in top_kw] or ["- No parseable copy found"])
    lines.extend(["","### Top 10 次级标签（Secondary）"])
    lines.extend([f"- {ph} ({cnt}) | competitor: {comp}" for ph,cnt,comp in top_ph] or ["- No parseable phrases found"])
    lines.extend(["","## 4) Google paid keywords 参考","","### Top 20 意图词"])
    lines.extend([f"- {kw} ({cnt})" for kw,cnt in g20] or ["- No paid keyword data returned"])
    lines.extend(["","### 投放最积极 Top 5 域名"])
    lines.extend([f"- {d}: {s}" for d,s in top_domains] or ["- No domain coverage data"])
    lines.extend(["","## 5) 对我们投放策略的直接启发","","1. Keep Meta objective split at campaign level: Click-to-Message first, Lead Form second, and isolate website traffic for retargeting only.","2. Standardize WhatsApp-first CTA variants in Cantonese + English for HK SMB restaurant owners and test by cuisine cluster.","3. Mirror top competitor intent terms into Google exact/phrase match lists, but route conversion to message-first landing flows.","4. Build ad creative sets around onboarding speed, monthly flexibility, and no hardware lock-in to counter common POS switching friction.","5. Use Meta manual audit queue weekly to capture new offer hooks and rotate copy templates every 14 days."])
    path.parent.mkdir(parents=True,exist_ok=True);path.write_text("\n".join(lines)+"\n",encoding="utf-8")

def extract_zip_if_needed(zip_path:Path,extract_dir:Path,logger:StepLogger)->Path:
    if not zip_path.exists():logger.log(f"Zip not found at {zip_path}, scanning repo only");return extract_dir
    extract_dir.mkdir(parents=True,exist_ok=True);marker=extract_dir/".extracted_ok";should=(not marker.exists()) or marker.stat().st_mtime<zip_path.stat().st_mtime
    if should:
        with zipfile.ZipFile(zip_path,"r") as zf:zf.extractall(extract_dir)
        marker.write_text(now_iso(),encoding="utf-8");logger.log(f"Extracted zip: {zip_path} -> {extract_dir}")
    else:logger.log(f"Using existing extracted directory: {extract_dir}")
    return extract_dir

def parse_args()->argparse.Namespace:
    p=argparse.ArgumentParser(description="Run HK POS ads intelligence pipeline")
    p.add_argument("--zip-path",default="../hk-pos-competitive-analysis.zip");p.add_argument("--extract-dir",default="input/extracted_hk_pos_competitive_analysis")
    p.add_argument("--database",default="hk");p.add_argument("--display-limit",type=int,default=200);p.add_argument("--dry-run-count",type=int,default=3);p.add_argument("--timeout-sec",type=int,default=20);p.add_argument("--skip-semrush",action="store_true")
    return p.parse_args()

def main()->int:
    a=parse_args()
    if a.display_limit>300:print("ERROR: --display-limit cannot exceed 300",file=sys.stderr);return 2
    base=Path(__file__).resolve().parent.parent;data_dir=base/"data";reports_dir=base/"reports";out_dir=base/"output"
    for d in (data_dir,reports_dir,out_dir):d.mkdir(exist_ok=True)
    logger=StepLogger(out_dir)
    stats={"extract_competitors":StepStat("extract_competitors"),"meta_collection":StepStat("meta_collection"),"semrush":StepStat("semrush"),"report":StepStat("report")}
    try:
        zip_path=(base/a.zip_path).resolve();extract_dir=(base/a.extract_dir).resolve()
        try:
            scan_root=extract_zip_if_needed(zip_path,extract_dir,logger)
        except OSError as exc:
            logger.log(f"WARN zip extraction failed: {exc}. Fallback to repository root.")
            scan_root=base
        if not scan_root.exists():scan_root=base
        try:
            competitors=extract_competitors_from_sources(scan_root,logger)
        except OSError as exc:
            logger.log(f"WARN source scan failed at {scan_root}: {exc}. Fallback to repository root scan.")
            competitors=extract_competitors_from_sources(base,logger)
        competitors=normalize_baseline_competitors(competitors)
        competitors=enrich_social_links(competitors,a.timeout_sec,logger)
        competitors=normalize_baseline_competitors(competitors)
        with_social=sum(1 for c in competitors if c.get("facebook_page_url") or c.get("instagram_handle"))
        coverage=(with_social/len(competitors)) if competitors else 0.0
        missing=[f"{c.get('competitor_name','')} ({c.get('website_domain','')})" for c in competitors if not (c.get("facebook_page_url") or c.get("instagram_handle"))]
        logger.log(f"Social URL coverage: {with_social}/{len(competitors)} = {coverage:.0%}")
        if missing:
            logger.log("Missing social URLs: " + "; ".join(missing))
        write_csv(data_dir/"competitors_master.csv",COMPETITOR_FIELDS,competitors);stats["extract_competitors"].success=len(competitors)
        logger.log(f"Quick reconnaissance: extracted_competitors={len(competitors)}, with_facebook_page_url={sum(1 for c in competitors if c.get('facebook_page_url'))}, planned_semrush_domains={len({normalize_domain(c.get('website_domain','') or c.get('website_url','')) for c in competitors if normalize_domain(c.get('website_domain','') or c.get('website_url',''))})}, display_limit={a.display_limit}")
        meta_rows,todos=collect_meta_ads(competitors,a.timeout_sec,logger)
        existing_todos=read_csv_rows(data_dir/"meta_ads_todo.csv")
        todos=merge_manual_todo_fields(todos,existing_todos)
        meta_rows=apply_manual_overrides(meta_rows,todos)
        todos=build_manual_todo_rows(meta_rows)
        todos=merge_manual_todo_fields(todos,existing_todos)
        write_csv(data_dir/"meta_ads_intel.csv",META_FIELDS,meta_rows)
        write_csv(data_dir/"meta_ads_todo.csv",META_FIELDS,todos)
        stats["meta_collection"].success=len([r for r in meta_rows if r.get("collection_method")=="web_auto" or r.get("status")=="active"])
        stats["meta_collection"].failed=len([r for r in meta_rows if r.get("status")!="active"])
        kw_rows=build_meta_keyword_rows(meta_rows);write_csv(data_dir/"meta_copy_keywords.csv",["row_type","competitor_name","ad_id_or_archive_id","ad_library_url","label_primary","label_secondary","label_reason","count"],kw_rows)
        sem_rows=[]
        existing_sem_rows=read_csv_rows(data_dir/"semrush_google_ads_signals.csv")
        if a.skip_semrush:
            logger.log("Semrush stage skipped by argument")
            sem_rows=existing_sem_rows
        else:
            key=os.getenv("SEMRUSH_API_KEY","").strip()
            if not key:
                logger.log("SEMRUSH_API_KEY is missing; keeping existing semrush output when available")
                sem_rows=existing_sem_rows or build_semrush_placeholder_rows(competitors,a.database,"missing_semrush_api_key")
            else:
                u_before,u_before_err=semrush_units_info(key)
                before_label=u_before if u_before else ("unavailable:" + (u_before_err or "unknown"))
                logger.log(f"Semrush units before run: {before_label}")
                sem_rows=collect_semrush_signals(competitors,key,a.database,a.display_limit,a.dry_run_count,logger)
                u_after,u_after_err=semrush_units_info(key)
                after_label=u_after if u_after else ("unavailable:" + (u_after_err or "unknown"))
                logger.log(f"Semrush units after run: {after_label}")
        write_csv(data_dir/"semrush_google_ads_signals.csv",SEMRUSH_FIELDS,sem_rows);stats["semrush"].success=len([r for r in sem_rows if r.get("paid_keywords_count")]);stats["semrush"].failed=len([r for r in sem_rows if not r.get("paid_keywords_count")])
        generate_report(reports_dir/"hk_competitor_ads_summary.md",competitors,meta_rows,kw_rows,sem_rows)
        build_ads_snapshot(base/"docs"/"data"/"ads_snapshot.json",meta_rows,kw_rows,sem_rows)
        stats["report"].success=1
        logger.log("Run summary:");[logger.log(f"- {s.step}: success={s.success}, failed={s.failed}") for s in stats.values()]
        reasons=Counter(r.get("error_reason","") for r in todos if r.get("error_reason"))
        if reasons:
            logger.log("Meta failure reason summary:")
            for reason,cnt in reasons.items():logger.log(f"  - {reason}: {cnt}")
        logger.flush();return 0
    except Exception as exc:
        logger.log(f"FATAL: {exc.__class__.__name__}: {sanitize_text(str(exc))}");logger.flush();return 1

if __name__=="__main__":raise SystemExit(main())
